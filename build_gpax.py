# -*- coding: utf-8 -*-
import re, os
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.utils import get_column_letter

SRC = "/home/user/Feypk/Plan_RU_Econ_Quant_v10.xlsx"
OUT = "/home/user/Feypk/GPAX_Calculator.xlsx"

# ---- read the 56 courses (term, code, name, credits) from v10 ----
src = load_workbook(SRC).active
courses=[]; term=None
for row in src.iter_rows(values_only=True):
    a=row[0]
    if isinstance(a,str):
        m=re.match(r'ปีที่\s*(\d+)\s*/\s*เทอม\s*(\d+)', a)
        if m: term="%s/%s"%(m.group(1),m.group(2))
    code=row[1]
    if isinstance(code,str) and re.match(r'^(ECO|MTH|STA|COS|SDM|FIN|ACC|RAM)\d',code):
        courses.append((term,code,row[2],int(row[3])))

# ---- styles ----
BLUE=PatternFill("solid",fgColor="4472C4"); NAVY=PatternFill("solid",fgColor="1F4E78")
PANEL=PatternFill("solid",fgColor="D9E1F2"); INPUT=PatternFill("solid",fgColor="FFF9D6")
GREEN=PatternFill("solid",fgColor="C6E0B4"); RED=PatternFill("solid",fgColor="FFC7CE"); ORG=PatternFill("solid",fgColor="F4B084")
GPAXC=PatternFill("solid",fgColor="2E75B6")
F=Font(name="Tahoma",size=10); FB=Font(name="Tahoma",size=10,bold=True)
FW=Font(name="Tahoma",size=10,bold=True,color="FFFFFF"); FT=Font(name="Tahoma",size=14,bold=True,color="FFFFFF")
FBIG=Font(name="Tahoma",size=22,bold=True,color="FFFFFF"); FNOTE=Font(name="Tahoma",size=9,italic=True)
thin=Side(style="thin",color="BFBFBF"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
C=Alignment(horizontal="center",vertical="center"); L=Alignment(horizontal="left",vertical="center",wrap_text=True)

wb=Workbook()
# ---- helper sheet: grade scale ----
sc=wb.active; sc.title="Scale"
sc["A1"]="เกรด"; sc["B1"]="แต้ม"
for i,(g,p) in enumerate([("A",4),("B+",3.5),("B",3),("C+",2.5),("C",2),("D+",1.5),("D",1),("F",0)],start=2):
    sc.cell(i,1,g); sc.cell(i,2,p)

ws=wb.create_sheet("GPAX"); ws.sheet_view.showGridLines=False
FIRST=10; NX=5; LAST=FIRST+len(courses)+NX-1   # courses + 5 blank rows
Dr="$D$%d:$D$%d"%(FIRST,LAST); Er="$E$%d:$E$%d"%(FIRST,LAST); Gr="$G$%d:$G$%d"%(FIRST,LAST); Ar="$A$%d:$A$%d"%(FIRST,LAST)

# Title + instructions
ws.merge_cells("A1:G1"); t=ws["A1"]; t.value="🎓 โปรแกรมคำนวณ GPAX อัตโนมัติ — เลือกเกรดในช่องสีเหลือง แล้วผลคำนวณเองทันที"
t.font=FT; t.fill=NAVY; t.alignment=C; ws.row_dimensions[1].height=26
ws.merge_cells("A2:G2"); ws["A2"]="เกณฑ์: A=4.0 | B+=3.5 | B=3.0 | C+=2.5 | C=2.0 | D+=1.5 | D=1.0 | F=0.0   •   GPAX = Σ(แต้ม×หน่วยกิต) ÷ Σ หน่วยกิตที่มีเกรด (F นับเป็น 0 แต่รวมหน่วยกิต)"
ws["A2"].font=FNOTE; ws["A2"].alignment=L; ws.row_dimensions[2].height=22

# Summary panel
ws.merge_cells("A4:G4"); s=ws["A4"]; s.value="📊 สรุปผล (คำนวณอัตโนมัติ)"; s.font=FW; s.fill=BLUE; s.alignment=L
def lab(cell,txt):
    ws[cell]=txt; ws[cell].font=FB; ws[cell].fill=PANEL; ws[cell].alignment=L; ws[cell].border=BD
ws.merge_cells("A5:B5"); lab("A5","GPAX สะสม")
ws.merge_cells("C5:D5"); ws["C5"]="=IF($G$5=0,\"\",SUM(%s)/$G$5)"%Gr
ws["C5"].font=FBIG; ws["C5"].fill=GPAXC; ws["C5"].alignment=C; ws["C5"].number_format="0.00"; ws["C5"].border=BD
ws.merge_cells("E5:F5"); lab("E5","หน่วยกิตที่มีเกรดแล้ว")
ws["G5"]="=SUMPRODUCT((%s<>\"\")*%s)"%(Er,Dr); ws["G5"].font=FB; ws["G5"].alignment=C; ws["G5"].border=BD; ws["G5"].number_format="0"
ws.row_dimensions[5].height=30
ws.merge_cells("A6:B6"); lab("A6","เทียบเกียรตินิยม*")
ws.merge_cells("C6:D6"); ws["C6"]="=IF($C$5=\"\",\"-\",IF($C$5>=3.5,\"อันดับ 1 🏆\",IF($C$5>=3.25,\"อันดับ 2 🥈\",\"ไม่เข้าเกณฑ์\")))"
ws["C6"].font=FB; ws["C6"].alignment=C; ws["C6"].border=BD
ws.merge_cells("E6:F6"); lab("E6","หน่วยกิตรวมทั้งหมด")
ws["G6"]="=SUM(%s)"%Dr; ws["G6"].font=FB; ws["G6"].alignment=C; ws["G6"].border=BD; ws["G6"].number_format="0"
ws.merge_cells("A7:B7"); lab("A7","🎯 ตั้งเป้า GPAX (กรอกตรงนี้)")
ws.merge_cells("C7:D7"); ws["C7"]=3.50; ws["C7"].font=FB; ws["C7"].fill=INPUT; ws["C7"].alignment=C; ws["C7"].border=BD; ws["C7"].number_format="0.00"
ws.merge_cells("E7:F7"); lab("E7","ต้องได้เฉลี่ยอีก (เต็ม 4)")
ws["G7"]="=IF(($G$6-$G$5)<=0,\"-\",IF($C$7=\"\",\"\",($C$7*$G$6-SUM(%s))/($G$6-$G$5)))"%Gr
ws["G7"].font=FB; ws["G7"].alignment=C; ws["G7"].border=BD; ws["G7"].number_format="0.00"
ws.merge_cells("A8:G8"); ws["A8"]="* เกณฑ์เกียรตินิยมแล้วแต่มหาลัย (ทั่วไป อันดับ1 ≥3.50, อันดับ2 ≥3.25) • ช่อง 'ต้องได้เฉลี่ยอีก': ถ้า >4 = เป็นไปไม่ได้ต้องลดเป้า, ถ้าติดลบ = เกินเป้าแล้ว"
ws["A8"].font=FNOTE; ws["A8"].alignment=L; ws["A8"].border=BD; ws.row_dimensions[8].height=22

# Course table header
hdr=["ปี/เทอม","รหัส","ชื่อวิชา","นก.","เกรด","แต้ม","แต้ม×นก."]
for j,h in enumerate(hdr,1):
    c=ws.cell(9,j,h); c.fill=BLUE; c.font=FW; c.alignment=C; c.border=BD
ws.freeze_panes="A10"

r=FIRST
for term,code,name,cr in courses:
    ws.cell(r,1,term).alignment=C
    ws.cell(r,2,code).alignment=C
    ws.cell(r,3,name)
    ws.cell(r,4,cr).alignment=C
    ws.cell(r,5).fill=INPUT; ws.cell(r,5).alignment=C   # grade input
    ws.cell(r,6,"=IFERROR(VLOOKUP($E%d,Scale!$A$2:$B$9,2,FALSE),\"\")"%r).alignment=C
    ws.cell(r,7,"=IF($E%d=\"\",0,$F%d*$D%d)"%(r,r,r)).alignment=C
    ws.cell(r,6).number_format="0.0"; ws.cell(r,7).number_format="0.0"
    for j in range(1,8): ws.cell(r,j).font=F; ws.cell(r,j).border=BD
    r+=1
# blank extra rows for user-added courses
for _ in range(NX):
    ws.cell(r,5).fill=INPUT; ws.cell(r,5).alignment=C
    ws.cell(r,6,"=IFERROR(VLOOKUP($E%d,Scale!$A$2:$B$9,2,FALSE),\"\")"%r).alignment=C
    ws.cell(r,7,"=IF($E%d=\"\",0,$F%d*$D%d)"%(r,r,r)).alignment=C
    ws.cell(r,6).number_format="0.0"; ws.cell(r,7).number_format="0.0"
    for j in range(1,8): ws.cell(r,j).font=F; ws.cell(r,j).border=BD
    r+=1
# totals row
ws.cell(r,3,"รวม").font=FB; ws.cell(r,3).alignment=Alignment(horizontal="right")
ws.cell(r,4,"=SUM(%s)"%Dr).font=FB; ws.cell(r,4).alignment=C; ws.cell(r,4).number_format="0"
ws.cell(r,7,"=SUM(%s)"%Gr).font=FB; ws.cell(r,7).alignment=C; ws.cell(r,7).number_format="0.0"
for j in range(1,8): ws.cell(r,j).fill=PANEL; ws.cell(r,j).border=BD

# ---- grade dropdown ----
dv=DataValidation(type="list",formula1='"A,B+,B,C+,C,D+,D,F"',allow_blank=True,showDropDown=False)
dv.prompt="เลือกเกรด"; dv.promptTitle="เกรด"
ws.add_data_validation(dv); dv.add("E%d:E%d"%(FIRST,LAST))

# ---- conditional formatting on grade column ----
gz="E%d:E%d"%(FIRST,LAST)
ws.conditional_formatting.add(gz,CellIsRule(operator="equal",formula=['"A"'],fill=GREEN))
ws.conditional_formatting.add(gz,CellIsRule(operator="equal",formula=['"B+"'],fill=GREEN))
ws.conditional_formatting.add(gz,CellIsRule(operator="equal",formula=['"D"'],fill=ORG))
ws.conditional_formatting.add(gz,CellIsRule(operator="equal",formula=['"D+"'],fill=ORG))
ws.conditional_formatting.add(gz,CellIsRule(operator="equal",formula=['"F"'],fill=RED))

# ---- per-term GPA table (right side) ----
ws.merge_cells("I4:K4"); h=ws["I4"]; h.value="GPA รายเทอม"; h.fill=BLUE; h.font=FW; h.alignment=C; h.border=BD
for j,t in enumerate(["เทอม","GPA","นก."],0):
    c=ws.cell(5+0, 9+j) # placeholder
terms=["1/1","1/2","2/1","2/2","3/1","3/2","4/1","4/2"]
ws.cell(5,9,"เทอม").font=FB; ws.cell(5,10,"GPA").font=FB; ws.cell(5,11,"นก.").font=FB
for j in (9,10,11): ws.cell(5,j).fill=PANEL; ws.cell(5,j).alignment=C; ws.cell(5,j).border=BD
for i,t in enumerate(terms):
    rr=6+i
    ws.cell(rr,9,t).alignment=C; ws.cell(rr,9).font=F
    cr_f="SUMPRODUCT((%s=$I%d)*(%s<>\"\")*%s)"%(Ar,rr,Er,Dr)
    q_f="SUMPRODUCT((%s=$I%d)*%s)"%(Ar,rr,Gr)
    ws.cell(rr,10,"=IF(%s=0,\"\",%s/%s)"%(cr_f,q_f,cr_f)).alignment=C; ws.cell(rr,10).number_format="0.00"; ws.cell(rr,10).font=FB
    ws.cell(rr,11,"=%s"%cr_f).alignment=C; ws.cell(rr,11).number_format="0"; ws.cell(rr,11).font=F
    for j in (9,10,11): ws.cell(rr,j).border=BD
ws.conditional_formatting.add("J6:J13",ColorScaleRule(start_type="num",start_value=0,start_color="F8696B",
    mid_type="num",mid_value=2,mid_color="FFEB84",end_type="num",end_value=4,end_color="63BE7B"))

# mini grade scale (right side, lower)
ws.cell(15,9,"เกณฑ์เกรด").font=FB; ws.cell(15,9).fill=PANEL
for i,(g,p) in enumerate([("A","4.0"),("B+","3.5"),("B","3.0"),("C+","2.5"),("C","2.0"),("D+","1.5"),("D","1.0"),("F","0.0")]):
    ws.cell(16+i,9,g).alignment=C; ws.cell(16+i,9).font=F; ws.cell(16+i,9).border=BD
    ws.cell(16+i,10,p).alignment=C; ws.cell(16+i,10).font=F; ws.cell(16+i,10).border=BD

for col,w in zip("ABCDEFG",[8,9,40,5,8,7,10]): ws.column_dimensions[col].width=w
for col,w in zip("IJK",[8,8,7]): ws.column_dimensions[col].width=w
try: wb.calculation.fullCalcOnLoad=True
except Exception: pass
wb.active=wb.sheetnames.index("GPAX")   # open on calculator, not the lookup
sc.sheet_state="hidden"
wb.save(OUT); print("SAVED",OUT,"courses=",len(courses),"rows",FIRST,"-",LAST)
