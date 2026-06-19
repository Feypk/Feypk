# -*- coding: utf-8 -*-
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/Feypk/Roadmap_CFA_Modeling_Internship.xlsx"

C_TITLE = PatternFill("solid", fgColor="1F4E78")
C_HEAD  = PatternFill("solid", fgColor="4472C4")
C_SEC   = PatternFill("solid", fgColor="2E75B6")
C_STAR  = PatternFill("solid", fgColor="FFF2CC")   # เน้น ปี3 ซัมเมอร์
C_CFA   = PatternFill("solid", fgColor="E2EFDA")
C_MOD   = PatternFill("solid", fgColor="DDEBF7")
C_INT   = PatternFill("solid", fgColor="FCE4D6")
C_QNT   = PatternFill("solid", fgColor="EDEDED")
F   = Font(name="Tahoma", size=10)
FB  = Font(name="Tahoma", size=10, bold=True)
FH  = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
FT  = Font(name="Tahoma", size=14, bold=True, color="FFFFFF")
FSEC= Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
FN  = Font(name="Tahoma", size=9, italic=True)
thin = Side(style="thin", color="BFBFBF")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
L = Alignment(horizontal="left", vertical="center", wrap_text=True)
LT= Alignment(horizontal="left", vertical="top", wrap_text=True)
Cc= Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook(); ws = wb.active; ws.title = "Roadmap"; ws.sheet_view.showGridLines = False
NC = 5
r = 1
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NC)
c=ws.cell(r,1,"Roadmap: CFA + Financial Modeling + Internship  (+ รักษาทางเลือก Quant)")
c.font=FT; c.fill=C_TITLE; c.alignment=Cc; ws.row_dimensions[r].height=26; r+=1
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NC)
c=ws.cell(r,1,"แผนเสริมนอกหลักสูตร คู่กับแผนเรียน v9 • เป้าหมาย: เข้าสาย Finance/Investment Banking โดยยังเปิดประตู Quant (MFE/MSc) ไว้ • ตัวตัดสินจริงของ IB = modeling + ฝึกงาน + เครือข่าย (มากกว่าเกรด/ปริญญา)")
c.font=FN; c.alignment=L; ws.row_dimensions[r].height=30; r+=2

def section(title):
    global r
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NC)
    c=ws.cell(r,1,title); c.font=FSEC; c.fill=C_SEC; c.alignment=L
    ws.row_dimensions[r].height=20; r+=1

def bullets(lines, fill=None):
    global r
    for ln in lines:
        ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NC)
        c=ws.cell(r,1,ln); c.font=F; c.alignment=LT; c.border=BD
        if fill: c.fill=fill
        ws.row_dimensions[r].height=16; r+=1

# ---------- Section A: Timeline matrix ----------
section("A) ไทม์ไลน์รวม (อ่านแนวนอน: แต่ละช่วงเวลาต้องทำอะไรในแต่ละสาย)")
heads=["ช่วงเวลา","🎓 CFA","📊 Financial Modeling","💼 Internship & Networking","🧮 Quant (รักษาทางเลือก)"]
for j,h in enumerate(heads,1):
    c=ws.cell(r,j,h); c.font=FH; c.fill=C_HEAD; c.alignment=Cc; c.border=BD
r+=1
rows=[
 ("ปี 1","เน้น GPA + พื้นฐาน","ฝึก Excel ให้คล่อง (shortcuts)","เข้าชมรมการเงิน/ลงทุน • สร้าง LinkedIn","ตั้งใจวิชาเลข/พิสูจน์ • เริ่ม Python"),
 ("ปี 1 ซัมเมอร์","–","Excel + เริ่ม 3-statement","อาสา/หาประสบการณ์ • ลงทุนจริงเล็กน้อย","โปรเจ็กต์ Python ลง GitHub"),
 ("ปี 2","อ่าน L1 เป็น background (เนื้อทับวิชา)","3-statement model + เริ่ม DCF","ฝึกงานสั้น: แบงก์/โบรกฯ/Big4/fintech","Probability/Stat • รักษาเกรดวิชาสัญญาณ"),
 ("ปี 2 ซัมเมอร์","–","DCF + Comparable (comps)","★ ฝึกงานครั้งแรก (การเงิน/บัญชี/วิจัย)","RA ผู้ช่วยอาจารย์ / Kaggle"),
 ("ปี 3","ลงทะเบียน CFA L1 + เริ่มติว","LBO + Merger model","เครือข่ายเชิงรุก (alumni/informational) • สมัครฝึกงานปี3","ทำ research paper • รักษาเกรด Real Analysis/Econometrics"),
 ("★ ปี 3 ซัมเมอร์","(เป้าหมายหลัก)","พอร์ตโมเดล 3–5 ชิ้นพร้อมโชว์","★★ ฝึกงาน IB/Finance ตัวจริง (Summer Analyst)","จุดตัดสิน: lean IB หรือ Quant • ถ้า Quant เริ่มเล็ง MFE+GRE"),
 ("ปี 4","สอบ CFA L1","ใช้โมเดลตอบ technical interview","แปลงฝึกงาน→offer / สมัคร full-time","ยื่นสมัคร MFE/MSc (ถ้าเลือกสาย Quant)"),
 ("หลังจบ","CFA L2 (ระหว่างทำงาน)","โมเดลขั้นสูงในงานจริง","เริ่มงาน analyst หรือเรียนต่อ","–"),
]
fills=[None,C_CFA,C_MOD,C_INT,C_QNT]
for row in rows:
    star = row[0].startswith("★")
    for j,v in enumerate(row,1):
        c=ws.cell(r,j,v); c.font=FB if (j==1 or star) else F; c.border=BD; c.alignment=LT
        fl = C_STAR if star else fills[j-1]
        if fl: c.fill = fl
    ws.row_dimensions[r].height=30; r+=1
r+=1

# ---------- Section B: CFA ----------
section("B) CFA — เหมาะกับสาย AM/ER/Research มากสุด • สำหรับ IB = สัญญาณเสริม • สำหรับ Quant = ไม่จำเป็น")
bullets([
 "• Level 1: ลงทะเบียนได้เมื่อเหลือ ≤ ~2 ปีก่อนจบ • สอบได้หลายรอบต่อปี (ก.พ./พ.ค./ส.ค./พ.ย.)",
 "• เนื้อ L1 ที่ 'แผนเรียน v9 ปูให้แล้ว': Economics (ECO) • Quant methods (STA) • Financial Reporting (ACC4209) • Corporate Finance (FIN3211) • Equity (FIN3209) • Fixed Income/Derivatives (ECO4327) • Portfolio (FIN3205)",
 "• ต้องติวเสริมเอง: Ethics & Professional Standards (ออกเยอะ) + Alternative Investments + การรวบยอดทั้งเล่ม",
 "• ลำดับแนะนำ: ปี2 อ่านคู่วิชา → ปี3 ลงทะเบียน+ติว → ปี4 สอบ L1 → หลังจบทำงานค่อยต่อ L2/L3",
], C_CFA)
r+=1

# ---------- Section C: Modeling ----------
section("C) Financial Modeling — เช็กลิสต์ทักษะ (หัวใจที่หลักสูตรให้ไม่ได้ ต้องฝึกเอง)")
bullets([
 "☐ Excel เร็ว ไม่ใช้เมาส์ (shortcuts, INDEX/MATCH, ตาราง, sensitivity)",
 "☐ 3-Statement Model (เชื่อมงบดุล–กำไรขาดทุน–กระแสเงินสด)   ← ต่อยอดจาก ACC4209",
 "☐ DCF Valuation (FCF, WACC, terminal value)   ← ต่อยอดจาก FIN3211/FIN3209",
 "☐ Comparable Companies (trading comps) + Precedent Transactions",
 "☐ LBO Model (leveraged buyout)",
 "☐ M&A / Merger Model (accretion / dilution)",
 "• แหล่งฝึก: Wall Street Prep, Breaking Into Wall Street (BIWS), CFI (FMVA cert), Macabacus, Mergers & Inquisitions (อ่านฟรี)",
 "• เป้าหมาย: มีพอร์ตโมเดล 3–5 ชิ้น (เช่น ปั้นโมเดลหุ้น SET จริง) ไว้โชว์ก่อนสัมภาษณ์ฝึกงาน",
], C_MOD)
r+=1

# ---------- Section D: Internship & Networking ----------
section("D) Internship & Networking — ★ ฝึกงานปี 3 ซัมเมอร์ = ตัวตัดสินที่สุด")
bullets([
 "• ที่ฝึกงานไทย: IBD/วาณิชธนกิจของแบงก์ (SCBX, KBank, BBL, Krungsri) • บล./โบรกเกอร์ (IB/Research) • Big4 Deal Advisory & Valuation (Deloitte/PwC/EY/KPMG) • บลจ. (Asset Management) • ตลาดหลักทรัพย์ (SET)/ก.ล.ต. • ธปท. (สาย Research)",
 "• เครือข่าย: LinkedIn (โปรไฟล์ + คอนเทนต์สม่ำเสมอ) • informational interview กับรุ่นพี่/ศิษย์เก่า • cold outreach อย่างสุภาพ",
 "• แข่งขัน (เร่งโปรไฟล์เร็ว): CFA Research Challenge • Valuation/Stock-Pitch competition • Trading competition • Case competition",
 "• เตรียมก่อนสมัคร: Resume 1 หน้า (เน้น quant + finance projects) • พอร์ตโมเดล • ฝึกตอบ technical (accounting, valuation, brain-teaser)",
 "• ลำดับ: ปี2 ฝึกงานสั้นสะสมประสบการณ์ → ปี3 ซัมเมอร์ ฝึกงาน IB/Finance ตัวจริง → ปี4 แปลงเป็น return offer",
], C_INT)
r+=1

# ---------- Section E: IB vs Quant ----------
section("E) จุดตัดสิน IB vs Quant — ตัดสินปลายปี 3 (แผน v9 รองรับทั้งสอง ยืดเวลาเลือกได้)")
bullets([
 "• ถ้า lean IB: ทุ่ม Financial Modeling + CFA + ฝึกงาน IB + networking เป็นหลัก",
 "• ถ้า lean Quant: ทุ่ม research paper + พอร์ต Python/C++ (GitHub) + GRE + สมัคร MFE/MSc • เกรด Real Analysis/Probability/Theory of Stats/Econometrics คือสัญญาณ",
 "• ข่าวดี: วิชาในแผน v9 ปูพื้นทั้งสองทางไว้แล้ว — ไม่ต้องรีบเลือกตั้งแต่ปี 1–2",
 "• หมายเหตุ: ทักษะ overlap เยอะ (programming, valuation, stats) → ทำคู่กันได้ถึงปี 3 แล้วค่อยโฟกัส",
], C_QNT)

for j,w in enumerate([16,26,30,34,30],1):
    ws.column_dimensions[get_column_letter(j)].width = w
wb.save(OUT); print("SAVED", OUT, "bytes=", os.path.getsize(OUT))
