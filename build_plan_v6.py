# -*- coding: utf-8 -*-
import base64
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/Feypk/Plan_RU_Econ_Quant_v6.xlsx"

# ---- category fills (ตาม legend เดิม) ----
FILL = {
    "econ":  PatternFill("solid", fgColor="DDEBF7"),  # ฟ้า = เศรษฐศาสตร์
    "quant": PatternFill("solid", fgColor="E2EFDA"),  # เขียว = คณิต-สถิติ-โปรแกรม
    "fin":   PatternFill("solid", fgColor="FCE4D6"),  # ส้ม = การเงิน (ECO/FIN)
    "gen":   PatternFill("solid", fgColor="F2F2F2"),  # เทา = ศึกษาทั่วไป
    "new":   PatternFill("solid", fgColor="FFF2CC"),  # เหลือง = วิชาเพิ่มใหม่ v6
}
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
TERM_FILL   = PatternFill("solid", fgColor="D9E1F2")
TITLE_FILL  = PatternFill("solid", fgColor="1F4E78")

FONT       = Font(name="Tahoma", size=10)
FONT_B     = Font(name="Tahoma", size=10, bold=True)
FONT_HDR   = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Tahoma", size=13, bold=True, color="FFFFFF")
FONT_NOTE  = Font(name="Tahoma", size=9, italic=True)

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center", wrap_text=True)

HEADERS = ["✅", "รหัส", "ชื่อวิชา", "นก.",
           "หมวด / เหตุผล", "บุพวิชา",
           "เทียบเท่า Top U / หมายเหตุ"]

# term -> (credit_total, rows[]) ; row = (code,name,credits,reason,prereq,note,cat)
PLAN = [
 ("ปีที่ 1 / เทอม 1", 21, [
   ("ECO1121","หลักเศรษฐศาสตร์จุลภาค",3,"แกน econ","-","Micro (MIT 14.01)","econ"),
   ("ECO1123","คณิต–สถิติเบื้องต้นสำหรับ ศ.",3,"แกน • ด่านทดสอบเลข","-","-","econ"),
   ("ACC1103","การบัญชีการเงิน",3,"แกน econ","-","-","econ"),
   ("MTH1101","Calculus & Analytic Geometry I",3,"เลือกเสรี • แคลคูลัสตัวจริง","-","Calculus I (MIT 18.01)","quant"),
   ("RAM1111","English in Daily Life",3,"ศึกษาทั่วไป 1.1 (อังกฤษ บังคับ)","-","-","gen"),
   ("RAM1103","Thai for Communication at Work",3,"ศึกษาทั่วไป 1.1 (ไทย)","-","-","gen"),
   ("COS1101","Introduction to Computer Science",3,"★ เพิ่มใหม่ v6 • เก็บเพิ่ม • พื้นฐาน CS ไม่มีบุพวิชา","-","ปลดล็อกสาย COS","new"),
 ]),
 ("ปีที่ 1 / เทอม 2", 21, [
   ("ECO1122","หลักเศรษฐศาสตร์มหภาค",3,"แกน econ","ECO1123","Macro (MIT 14.02)","econ"),
   ("MTH1102","Calculus & Analytic Geometry II",3,"เลือกเสรี • แคลคูลัสต่อเนื่อง","MTH1101","Calculus II (MIT 18.01-02)","quant"),
   ("MTH2001","Fundamental Concepts in Math 1",3,"★ วิชาเขียนพิสูจน์ (ฐานก่อน Real Analysis)","-","Proof course (Harvard Math 101)","quant"),
   ("MTH2206","Matrix Theory & Linear Algebra 1",3,"⬅ ย้ายมาจากปี2/เทอม1 • โทคณิต (PR แค่ MTH1101)","MTH1101","Linear Algebra (MIT 18.06)","quant"),
   ("RAM1131","Digital Literacy",3,"ศึกษาทั่วไป 1.1 (เทคโนโลยี)","-","-","gen"),
   ("RAM1301","คุณธรรมคู่ความรู้",3,"ศึกษาทั่วไป 1.3 (พลเมือง บังคับ)","-","-","gen"),
   ("COS1103","Algorithms and Programming",3,"★ เพิ่มใหม่ v6 • ต่อจาก COS1101 • ปลดล็อก COS2102 (C++)","COS1101","เลิกต้องลุ้นเทียบ Python","new"),
 ]),
 ("ปีที่ 2 / เทอม 1", 21, [
   ("ECO2121","ทฤษฎีจุลภาค 1",3,"แกน econ","ECO1121,1123","Intermediate Micro","econ"),
   ("ECO2122","ทฤษฎีมหภาค 1",3,"แกน econ","ECO1122,1123","Intermediate Macro","econ"),
   ("ECO2125","คณิตศาสตร์สำหรับเศรษฐศาสตร์",3,"แกน • วิชาสัญญาณ","ECO1123","Math for Econ","econ"),
   ("ECO2129","เศรษฐกิจประเทศไทย",3,"⬅ เลื่อนมาจากปี1/เทอม2 • แกน econ (เรียนหลัง ECO1122)","ECO1121,1122","-","econ"),
   ("MTH2103","Calculus & Analytic Geometry III",3,"โทคณิต • แคลคูลัสหลายตัวแปร","MTH1102","Multivariable Calc (Harvard 21a)","quant"),
   ("STA3101","Probability Theory 1",3,"★ ฐานความน่าจะเป็น (ปลดล็อก Prob 2)","MTH1102","Probability (Harvard Stat 110)","quant"),
   ("RAM1112","ภาษาและวัฒนธรรมอังกฤษ",3,"ศึกษาทั่วไป 1.1 (ภาษาต่างประเทศ)","-","-","gen"),
 ]),
 ("ปีที่ 2 / เทอม 2", 21, [
   ("ECO2123","ทฤษฎีจุลภาค 2",3,"แกน econ","ECO2121","-","econ"),
   ("ECO2124","ทฤษฎีมหภาค 2",3,"แกน econ","ECO2122","-","econ"),
   ("ECO2126","สถิติสำหรับเศรษฐศาสตร์",3,"แกน • วิชาสัญญาณ","(เช็กภาควิชา: STA2003 หรือเทียบ ECO1123)","Statistics (MIT 14.30)","econ"),
   ("MTH2104","Advanced Calculus 1",3,"โทคณิต • ปลดล็อก Real Analysis","MTH2103","Advanced Calculus","quant"),
   ("MTH3204","Matrix Theory & Linear Algebra 2",3,"โทคณิต • Linear Algebra ขั้นสูง","MTH2206","Linear Algebra II","quant"),
   ("RAM1141","ศาสตร์แห่งบุคลิกภาพ",3,"ศึกษาทั่วไป 1.1 (ปรับตัว)","-","-","gen"),
   ("RAM1204","Mathematics & Statistics for Daily Life",3,"ศึกษาทั่วไป 1.2 (ทักษะการคิด)","-","-","gen"),
 ]),
 ("ปีที่ 3 / เทอม 1", 21, [
   ("ECO3121","เศรษฐศาสตร์จุลภาคระดับสูง",3,"เชี่ยวชาญ (ทฤษฎีและเชิงปริมาณ)","ECO2123,2125","Advanced Micro","econ"),
   ("ECO3123","เศรษฐมิติเบื้องต้น",3,"เชี่ยวชาญ • วิชาดาวเด่น","ECO2125,2126","Econometrics (MIT 14.32)","econ"),
   ("ECO3320","เศรษฐศาสตร์การเงินและการธนาคาร",3,"เลือกกลุ่มอื่น–การเงิน","ECO2124","-","fin"),
   ("MTH2106","Differential Equations",3,"โทคณิต • สมการเชิงอนุพันธ์","MTH1102","Diff Equations (MIT 18.03)","quant"),
   ("STA3102","Probability Theory 2",3,"★ ปลดล็อก Theory of Statistics","STA3101+MTH2103","Probability II","quant"),
   ("ECO2130","English for Economics",3,"แกน econ","-","-","econ"),
   ("STA3301","Regression Analysis",3,"★ สถิติประยุกต์เสริม econometrics","พื้นสถิติ (ECO2126)","Applied Regression","quant"),
 ]),
 ("ปีที่ 3 / เทอม 2", 21, [
   ("ECO3122","เศรษฐศาสตร์มหภาคระดับสูง",3,"เชี่ยวชาญ","ECO2124,2125","Advanced Macro","econ"),
   ("ECO3323","การบริหารความเสี่ยงองค์กร",3,"เลือกกลุ่มอื่น–การเงิน","ECO2121","-","fin"),
   ("ECO2128","ระเบียบวิธีวิจัยทางเศรษฐศาสตร์",3,"แกน econ","ECO2123-2126","-","econ"),
   ("MTH4404","Introduction to Real Analysis 1",3,"เก็บเพิ่ม • grad-track","MTH2104(+MTH2001)","Real Analysis (MIT 18.100)","quant"),
   ("STA4101","Theory of Statistics 1",3,"★ Mathematical Statistics (MFE core)","STA3102","Statistical Inference","quant"),
   ("SDM2102","Foundations of Python Programming",3,"เก็บเพิ่ม • Python มีหน่วยกิต","-","Programming (MIT 6.100A)","quant"),
   ("RAM1312","Contemporary Culture & Digital Disruption",3,"ศึกษาทั่วไป 1.3 (ศิลปะ)","-","-","gen"),
 ]),
 ("ปีที่ 4 / เทอม 1", 21, [
   ("ECO4124","Time-Series Econometrics",3,"เชี่ยวชาญ (เลือก)","ECO3123","Time-Series Econometrics","econ"),
   ("ECO4125","Microeconometrics",3,"เชี่ยวชาญ (เลือก)","ECO3123","-","econ"),
   ("ECO4326","Mathematical Finance",3,"เลือกกลุ่มอื่น–การเงิน","ECO3323","Mathematical Finance","fin"),
   ("MTH3407","Numerical Analysis 1",3,"เก็บเพิ่ม • computational","MTH2206+MTH2104","Numerical Analysis","quant"),
   ("MTH3611","Mathematical Stochastic Models",3,"★ Stochastic Processes (quant finance)","MTH2106","บุพวิชาสายคณิต (ไม่ใช่ STA2003)","quant"),
   ("FIN3319","Fintech and Digital Assets",3,"★ FIN เบา • algo/digital assets","FIN2101*","*ขอเทียบพื้นการเงิน","fin"),
   ("RAM1302","Politics and Law in Daily Life",3,"ศึกษาทั่วไป 1.3 (พลเมือง เลือก)","-","-","gen"),
 ]),
 ("ปีที่ 4 / เทอม 2", 21, [
   ("ECO4123","Game Theory (ทฤษฎีเกม)",3,"เชี่ยวชาญ (เลือก)","ECO2123,2125","Game Theory (MIT 14.12)","econ"),
   ("ECO4327","Financial Risk Mgmt & Derivatives",3,"เลือกกลุ่มอื่น–การเงิน","ECO3320,3323","-","fin"),
   ("MTH4405","Introduction to Real Analysis 2",3,"★ สำหรับกลุ่ม Reach (Oxford/Imperial)","MTH4404","Real Analysis II","quant"),
   ("COS2102","Object-Oriented Programming (C++)",3,"★ programming ลึก (quant dev)","COS1103","✅ ปลดล็อกจริงจาก COS1103 แล้ว","quant"),
   ("FIN3209","Securities Analysis",3,"★ FIN • valuation ภาคปฏิบัติ","FIN2101*","*ขอเทียบพื้นการเงิน","fin"),
   ("FIN3205","Principles & Policies of Investment",3,"★ FIN • ฐานการลงทุน/พอร์ต","FIN2101*","*ขอเทียบพื้นการเงิน","fin"),
   ("RAM1211","Principles of Entrepreneurship",3,"ศึกษาทั่วไป 1.2 (ผู้ประกอบการ)","-","-","gen"),
 ]),
]

wb = Workbook()
ws = wb.active
ws.title = "แผนรายเทอม v6"
ws.sheet_view.showGridLines = False

NC = 7
# Title
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC)
c = ws.cell(1,1,"แผนการเรียน ป.ตรี เศรษฐศาสตร์ ม.รามคำแหง — สาย Quant (v6: 168 นก.)")
c.font = FONT_TITLE; c.fill = TITLE_FILL; c.alignment = CENTER
ws.row_dimensions[1].height = 24
# Note
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NC)
note = ("แก้จาก v5: เติม COS1101 (ปี1/เทอม1) + COS1103 (ปี1/เทอม2) อย่างละ 3 นก. "
        "→ ปี1 เต็ม 21 นก./เทอม • รวมทั้งโปรแกรม = 168 นก. (หลักสูตรแกน 120 + เก็บเพิ่ม 48) "
        "• COS1101→COS1103 ปลดล็อก COS2102 (C++) ในปี4 ได้จริง")
c = ws.cell(2,1,note); c.font = FONT_NOTE; c.alignment = LEFT
ws.row_dimensions[2].height = 42
# Legend
ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=NC)
c = ws.cell(3,1,"ฟ้า=เศรษฐศาสตร์ | เขียว=คณิต–สถิติ–โปรแกรม | ส้ม=การเงิน (ECO/FIN) | เทา=ศึกษาทั่วไป | เหลือง=วิชาเพิ่มใหม่ v6 | ⬅=ย้ายใน v5")
c.font = FONT_NOTE; c.alignment = LEFT

r = 5
# header
for j,h in enumerate(HEADERS, start=1):
    cell = ws.cell(r,j,h); cell.fill = HEADER_FILL; cell.font = FONT_HDR
    cell.alignment = CENTER if j in (1,2,4) else LEFT; cell.border = BORDER
ws.freeze_panes = ws.cell(r+1,1)
r += 1

for term, total, rows in PLAN:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC-1)
    tc = ws.cell(r,1,term); tc.fill = TERM_FILL; tc.font = FONT_B; tc.alignment = LEFT; tc.border = BORDER
    for j in range(2,NC):
        ws.cell(r,j).fill = TERM_FILL; ws.cell(r,j).border = BORDER
    cc = ws.cell(r,NC,"%d นก." % total); cc.fill = TERM_FILL; cc.font = FONT_B; cc.alignment = CENTER; cc.border = BORDER
    r += 1
    for code,name,cr,reason,pre,note2,cat in rows:
        vals = ["", code, name, cr, reason, pre, note2]
        for j,v in enumerate(vals, start=1):
            cell = ws.cell(r,j,v); cell.fill = FILL[cat]; cell.font = FONT; cell.border = BORDER
            cell.alignment = CENTER if j in (1,2,4) else LEFT
        r += 1

# summary
r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
c = ws.cell(r,1,"สรุปตรวจหน่วยกิต (โครงสร้างหลัก 120 ไม่เปลี่ยน)")
c.font = FONT_B; c.fill = TERM_FILL; c.alignment = LEFT; c.border = BORDER
r += 1
summ = [
 ("หมวด","ต้องการ","ในแผน","สถานะ"),
 ("ศึกษาทั่วไป","30","30","ครบ"),
 ("วิชาแกน","39","39","ครบ"),
 ("เชี่ยวชาญ (ทฤษฎีฯ)","18","18","ครบ"),
 ("เลือกกลุ่มอื่น (การเงิน)","12","12","ครบ"),
 ("วิชาโท (คณิตศาสตร์)","15","15","ครบ"),
 ("เลือกเสรี (MTH1101+MTH1102)","6","6","ครบ"),
 ("รวมหลักสูตร","120","120","ครบพอดี"),
 ("เก็บเพิ่ม (STEM + FIN + CS)","-","48","เสริมโปรไฟล์ quant"),
 ("รวมทั้งโปรแกรม","-","168","ทุกเทอม 21 นก. (เพดาน 22)"),
]
for i,(a,b,cc2,d) in enumerate(summ):
    fnt = FONT_B if i==0 else FONT
    fil = TERM_FILL if i==0 else None
    for j,v in enumerate([a,b,cc2,d], start=1):
        cell = ws.cell(r,j,v); cell.font = fnt; cell.border = BORDER
        if fil: cell.fill = fil
        cell.alignment = CENTER if (j in (2,3) and i>0) else LEFT
    r += 1

widths = [4,10,42,7,40,24,30]
for j,w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(j)].width = w

wb.save(OUT)

import os
size = os.path.getsize(OUT)
with open(OUT,"rb") as f:
    b64 = base64.b64encode(f.read()).decode()
with open(OUT + ".b64","w") as f:
    f.write(b64)
print("SAVED", OUT, "bytes=", size, "b64len=", len(b64))
