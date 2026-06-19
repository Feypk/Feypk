# -*- coding: utf-8 -*-
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/Feypk/Plan_RU_Econ_Quant_v9.xlsx"

FILL = {
    "econ":  PatternFill("solid", fgColor="DDEBF7"),
    "quant": PatternFill("solid", fgColor="E2EFDA"),
    "fin":   PatternFill("solid", fgColor="FCE4D6"),
    "gen":   PatternFill("solid", fgColor="F2F2F2"),
    "new":   PatternFill("solid", fgColor="FFF2CC"),
}
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
TERM_FILL   = PatternFill("solid", fgColor="D9E1F2")
TITLE_FILL  = PatternFill("solid", fgColor="1F4E78")
FONT       = Font(name="Tahoma", size=10)
FONT_B     = Font(name="Tahoma", size=10, bold=True)
FONT_HDR   = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Tahoma", size=13, bold=True, color="FFFFFF")
FONT_NOTE  = Font(name="Tahoma", size=9, italic=True)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center", wrap_text=True)
HEADERS = ["✅", "รหัส", "ชื่อวิชา", "นก.", "หมวด / เหตุผล", "บุพวิชา", "เทียบเท่า Top U / หมายเหตุ"]

PLAN = [
 ("ปีที่ 1 / เทอม 1", 21, [
   ("ECO1121","หลักเศรษฐศาสตร์จุลภาค",3,"แกน econ","-","Micro (MIT 14.01)","econ"),
   ("ECO1123","คณิต–สถิติเบื้องต้นสำหรับ ศ.",3,"แกน • ด่านทดสอบเลข","-","-","econ"),
   ("ACC1103","การบัญชีการเงิน",3,"แกน econ","-","พื้นบัญชี → ACC2202/FIN2101","econ"),
   ("MTH1101","Calculus & Analytic Geometry I",3,"เลือกเสรี • แคลคูลัสตัวจริง","-","Calculus I (MIT 18.01)","quant"),
   ("RAM1111","English in Daily Life",3,"ศึกษาทั่วไป 1.1 (อังกฤษ บังคับ)","-","-","gen"),
   ("RAM1103","Thai for Communication at Work",3,"ศึกษาทั่วไป 1.1 (ไทย)","-","-","gen"),
   ("COS1101","Introduction to Computer Science",3,"เก็บเพิ่ม • พื้นฐาน CS","-","ปลดล็อกสาย COS","quant"),
 ]),
 ("ปีที่ 1 / เทอม 2", 21, [
   ("ECO1122","หลักเศรษฐศาสตร์มหภาค",3,"แกน econ","ECO1123","Macro (MIT 14.02)","econ"),
   ("MTH1102","Calculus & Analytic Geometry II",3,"เลือกเสรี • แคลคูลัสต่อเนื่อง","MTH1101","Calculus II","quant"),
   ("MTH2001","Fundamental Concepts in Math 1",3,"★ พิสูจน์ (ฐานก่อน Real Analysis)","-","Proof course","quant"),
   ("MTH2206","Matrix Theory & Linear Algebra 1",3,"โทคณิต (PR แค่ MTH1101)","MTH1101","Linear Algebra (MIT 18.06)","quant"),
   ("RAM1131","Digital Literacy",3,"ศึกษาทั่วไป 1.1 (เทคโนโลยี)","-","-","gen"),
   ("RAM1301","คุณธรรมคู่ความรู้",3,"ศึกษาทั่วไป 1.3 (พลเมือง บังคับ)","-","-","gen"),
   ("COS1103","Algorithms and Programming",3,"เก็บเพิ่ม • ปลดล็อก COS2102","COS1101","-","quant"),
 ]),
 ("ปีที่ 2 / เทอม 1", 21, [
   ("ECO2121","ทฤษฎีจุลภาค 1",3,"แกน econ","ECO1121,1123","Intermediate Micro","econ"),
   ("ECO2122","ทฤษฎีมหภาค 1",3,"แกน econ","ECO1122,1123","Intermediate Macro","econ"),
   ("ECO2125","คณิตศาสตร์สำหรับเศรษฐศาสตร์",3,"แกน • วิชาสัญญาณ","ECO1123","Math for Econ","econ"),
   ("ECO2129","เศรษฐกิจประเทศไทย",3,"แกน econ","ECO1121,1122","-","econ"),
   ("MTH2103","Calculus & Analytic Geometry III",3,"โทคณิต • หลายตัวแปร","MTH1102","Multivariable Calc","quant"),
   ("STA3101","Probability Theory 1",3,"★ ฐานความน่าจะเป็น","MTH1102","Probability (Stat 110)","quant"),
   ("RAM1112","ภาษาและวัฒนธรรมอังกฤษ",3,"ศึกษาทั่วไป 1.1 (ภาษาต่างประเทศ)","-","-","gen"),
 ]),
 ("ปีที่ 2 / เทอม 2", 21, [
   ("ECO2123","ทฤษฎีจุลภาค 2",3,"แกน econ","ECO2121","-","econ"),
   ("ECO2124","ทฤษฎีมหภาค 2",3,"แกน econ","ECO2122","-","econ"),
   ("ECO2126","สถิติสำหรับเศรษฐศาสตร์",3,"แกน • วิชาสัญญาณ","(เช็กภาควิชา)","Statistics (MIT 14.30)","econ"),
   ("MTH2104","Advanced Calculus 1",3,"โทคณิต • ปลดล็อก Real Analysis","MTH2103","Advanced Calculus","quant"),
   ("MTH3204","Matrix Theory & Linear Algebra 2",3,"โทคณิต • ขั้นสูง","MTH2206","Linear Algebra II","quant"),
   ("RAM1141","ศาสตร์แห่งบุคลิกภาพ",3,"ศึกษาทั่วไป 1.1 (ปรับตัว)","-","-","gen"),
   ("ACC2202","Intermediate Accounting",3,"⬅ ขยับขึ้น v9 • บัญชีเชิงลึก (อ่านงบก่อนฝึกงาน)","ACC1103 (เช็กภาควิชา)","ฐานก่อน ACC4209","fin"),
 ]),
 ("ปีที่ 3 / เทอม 1", 21, [
   ("ECO3121","เศรษฐศาสตร์จุลภาคระดับสูง",3,"เชี่ยวชาญ (ทฤษฎีและเชิงปริมาณ)","ECO2123,2125","Advanced Micro","econ"),
   ("ECO3123","เศรษฐมิติเบื้องต้น",3,"เชี่ยวชาญ • วิชาดาวเด่น","ECO2125,2126","Econometrics (MIT 14.32)","econ"),
   ("ECO3320","เศรษฐศาสตร์การเงินและการธนาคาร",3,"เลือกกลุ่มอื่น–การเงิน","ECO2124","-","fin"),
   ("MTH2106","Differential Equations",3,"โทคณิต • สมการเชิงอนุพันธ์","MTH1102","Diff Equations (MIT 18.03)","quant"),
   ("STA3102","Probability Theory 2",3,"★ ปลดล็อก Theory of Statistics","STA3101+MTH2103","Probability II","quant"),
   ("ECO2130","English for Economics",3,"แกน econ","-","-","econ"),
   ("FIN2101","Business Finance (การเงินธุรกิจ)",3,"เก็บเพิ่ม • บุพวิชาของ FIN3xxx","ACC1103 (เช็กภาควิชา)","ฐานก่อน FIN3211/3209/3205","fin"),
 ]),
 ("ปีที่ 3 / เทอม 2", 21, [
   ("ECO3122","เศรษฐศาสตร์มหภาคระดับสูง",3,"เชี่ยวชาญ","ECO2124,2125","Advanced Macro","econ"),
   ("ECO3323","การบริหารความเสี่ยงองค์กร",3,"เลือกกลุ่มอื่น–การเงิน","ECO2121","-","fin"),
   ("ECO2128","ระเบียบวิธีวิจัยทางเศรษฐศาสตร์",3,"แกน econ • ต่อยอดเป็น research paper","ECO2123-2126","-","econ"),
   ("MTH4404","Introduction to Real Analysis 1",3,"เก็บเพิ่ม • สัญญาณ #1 grad school","MTH2104(+MTH2001)","Real Analysis (MIT 18.100)","quant"),
   ("STA4101","Theory of Statistics 1",3,"★ Mathematical Statistics (MFE core)","STA3102","Statistical Inference","quant"),
   ("ACC4209","Financial Reporting & Analysis",3,"⬅ ขยับขึ้น v9 • วิเคราะห์งบ (core IB — ก่อนฝึกงาน)","ACC2202 (เช็กภาควิชา)","financial statement analysis","fin"),
   ("SDM2102","Foundations of Python Programming",3,"เก็บเพิ่ม • Python","-","Programming (MIT 6.100A)","quant"),
 ]),
 ("ปีที่ 4 / เทอม 1", 21, [
   ("ECO4124","Time-Series Econometrics",3,"เชี่ยวชาญ (เลือก)","ECO3123","Time-Series Econometrics","econ"),
   ("ECO4125","Microeconometrics",3,"เชี่ยวชาญ (เลือก) • เช็ก causal inference","ECO3123","IV/DiD/RDD?","econ"),
   ("ECO4326","Mathematical Finance",3,"เลือกกลุ่มอื่น–การเงิน","ECO3323","Mathematical Finance","fin"),
   ("MTH3611","Mathematical Stochastic Models",3,"เก็บเพิ่ม • Stochastic (quant finance)","MTH2106","หัวใจ quant finance","quant"),
   ("FIN3211","Financial Management (Corporate Finance)",3,"เก็บเพิ่ม • การเงินองค์กร — เนื้อ IB","FIN2101 (อาจต้อง FIN3208)","capital structure/budgeting","fin"),
   ("RAM1204","Mathematics & Statistics for Daily Life",3,"⬅ ย้ายมาจากปี2/เทอม2 • ศึกษาทั่วไป 1.2","-","-","gen"),
   ("RAM1302","Politics and Law in Daily Life",3,"ศึกษาทั่วไป 1.3 (พลเมือง เลือก)","-","-","gen"),
 ]),
 ("ปีที่ 4 / เทอม 2", 21, [
   ("ECO4123","Game Theory (ทฤษฎีเกม)",3,"เชี่ยวชาญ (เลือก)","ECO2123,2125","Game Theory (MIT 14.12)","econ"),
   ("ECO4327","Financial Risk Mgmt & Derivatives",3,"เลือกกลุ่มอื่น–การเงิน","ECO3320,3323","ครอบคลุม derivatives","fin"),
   ("COS2102","Object-Oriented Programming (C++)",3,"เก็บเพิ่ม • C++ (quant dev)","COS1103","✅ ปลดล็อกจาก COS1103","quant"),
   ("FIN3209","Securities Analysis",3,"เก็บเพิ่ม • valuation ภาคปฏิบัติ","FIN2101","equity valuation","fin"),
   ("FIN3205","Principles & Policies of Investment",3,"เก็บเพิ่ม • ฐานการลงทุน/พอร์ต","FIN2101","investments","fin"),
   ("RAM1312","Contemporary Culture & Digital Disruption",3,"⬅ ย้ายมาจากปี3/เทอม2 • ศึกษาทั่วไป 1.3","-","-","gen"),
   ("RAM1211","Principles of Entrepreneurship",3,"ศึกษาทั่วไป 1.2 (ผู้ประกอบการ)","-","-","gen"),
 ]),
]

wb = Workbook(); ws = wb.active
ws.title = "แผนรายเทอม v9"; ws.sheet_view.showGridLines = False
NC = 7
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC)
c = ws.cell(1,1,"แผนการเรียน ป.ตรี เศรษฐศาสตร์ ม.รามคำแหง — Quant + Finance/IB (v9: 168 นก.)")
c.font = FONT_TITLE; c.fill = TITLE_FILL; c.alignment = CENTER; ws.row_dimensions[1].height = 24
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NC)
note = ("แก้จาก v8: ขยับ ACC2202 → ปี2/เทอม2 และ ACC4209 → ปี3/เทอม2 ให้เรียนจบสายบัญชี/อ่านงบ "
        "ก่อนฝึกงานปี3 ซัมเมอร์ • สลับ RAM1204, RAM1312 ไปปี4 • หน่วยกิต/โครงสร้างเท่าเดิม 168 นก.")
c = ws.cell(2,1,note); c.font = FONT_NOTE; c.alignment = LEFT; ws.row_dimensions[2].height = 42
ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=NC)
c = ws.cell(3,1,"ฟ้า=เศรษฐศาสตร์ | เขียว=คณิต–สถิติ–โปรแกรม | ส้ม=การเงิน/บัญชี | เทา=ศึกษาทั่วไป | ⬅=วิชาที่ขยับเทอมใน v9")
c.font = FONT_NOTE; c.alignment = LEFT
r = 5
for j,h in enumerate(HEADERS, start=1):
    cell = ws.cell(r,j,h); cell.fill = HEADER_FILL; cell.font = FONT_HDR
    cell.alignment = CENTER if j in (1,2,4) else LEFT; cell.border = BORDER
ws.freeze_panes = ws.cell(r+1,1); r += 1
for term, total, rows in PLAN:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC-1)
    tc = ws.cell(r,1,term); tc.fill = TERM_FILL; tc.font = FONT_B; tc.alignment = LEFT; tc.border = BORDER
    for j in range(2,NC): ws.cell(r,j).fill = TERM_FILL; ws.cell(r,j).border = BORDER
    cc = ws.cell(r,NC,"%d นก." % total); cc.fill = TERM_FILL; cc.font = FONT_B; cc.alignment = CENTER; cc.border = BORDER
    r += 1
    for code,name,cr,reason,pre,note2,cat in rows:
        for j,v in enumerate(["", code, name, cr, reason, pre, note2], start=1):
            cell = ws.cell(r,j,v); cell.fill = FILL[cat]; cell.font = FONT; cell.border = BORDER
            cell.alignment = CENTER if j in (1,2,4) else LEFT
        r += 1
r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
c = ws.cell(r,1,"สรุปตรวจหน่วยกิต (โครงสร้างหลัก 120 ไม่เปลี่ยน) • บัญชีจบก่อนฝึกงานปี3 ซัมเมอร์ ✅")
c.font = FONT_B; c.fill = TERM_FILL; c.alignment = LEFT; c.border = BORDER; r += 1
for i,(a,b,cc2,d) in enumerate([
 ("หมวด","ต้องการ","ในแผน","สถานะ"),("ศึกษาทั่วไป","30","30","ครบ"),("วิชาแกน","39","39","ครบ"),
 ("เชี่ยวชาญ (ทฤษฎีฯ)","18","18","ครบ"),("เลือกกลุ่มอื่น (การเงิน)","12","12","ครบ"),
 ("วิชาโท (คณิตศาสตร์)","15","15","ครบ"),("เลือกเสรี (MTH1101+1102)","6","6","ครบ"),
 ("รวมหลักสูตร","120","120","ครบพอดี"),("เก็บเพิ่ม (STEM+CS+Finance+ACC)","-","48","Quant+IB"),
 ("รวมทั้งโปรแกรม","-","168","ทุกเทอม 21 นก.")]):
    for j,v in enumerate([a,b,cc2,d], start=1):
        cell = ws.cell(r,j,v); cell.font = FONT_B if i==0 else FONT; cell.border = BORDER
        if i==0: cell.fill = TERM_FILL
        cell.alignment = CENTER if (j in (2,3) and i>0) else LEFT
    r += 1
for j,w in enumerate([4,10,42,7,42,26,28], start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
wb.save(OUT); print("SAVED", OUT, "bytes=", os.path.getsize(OUT))
