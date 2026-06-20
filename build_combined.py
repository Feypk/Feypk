# -*- coding: utf-8 -*-
import re, os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.utils import get_column_letter

OUT="/home/user/Feypk/Plan_RU_Econ_Quant_v11.xlsx"

# ---------- shared styles ----------
def P(c): return PatternFill("solid",fgColor=c)
BLUE=P("4472C4"); NAVY=P("1F4E78"); PANEL=P("D9E1F2"); SEC=P("2E75B6"); INPUT=P("FFF9D6")
FILL={"econ":P("DDEBF7"),"quant":P("E2EFDA"),"fin":P("FCE4D6"),"gen":P("F2F2F2")}
D_RED=P("FFC7CE"); D_ORG=P("F4B084"); D_YEL=P("FFE699"); D_GRN=P("C6E0B4"); GPAXC=P("2E75B6")
RISK={"low":P("C6E0B4"),"med":P("FFE699"),"high":P("F4B084"),"vhigh":P("FFC7CE")}
def dfill(d): return D_RED if d>=9 else D_ORG if d>=7 else D_YEL if d>=5 else D_GRN
F=Font(name="Tahoma",size=10); FB=Font(name="Tahoma",size=10,bold=True)
FW=Font(name="Tahoma",size=10,bold=True,color="FFFFFF"); FT=Font(name="Tahoma",size=13,bold=True,color="FFFFFF")
FSEC=Font(name="Tahoma",size=11,bold=True,color="FFFFFF"); FBIG=Font(name="Tahoma",size=20,bold=True,color="FFFFFF")
FNOTE=Font(name="Tahoma",size=9,italic=True)
thin=Side(style="thin",color="BFBFBF"); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
C=Alignment(horizontal="center",vertical="center",wrap_text=True); L=Alignment(horizontal="left",vertical="center",wrap_text=True)
LT=Alignment(horizontal="left",vertical="top",wrap_text=True)

# ---------- PLAN data (v10) : (code,name,cr,diff,reason,prereq,status,cat) ----------
PLAN=[
 ("ปีที่ 1 / เทอม 1",[("ECO1121","หลักเศรษฐศาสตร์จุลภาค",3,3,"แกน econ","-","✓","econ"),("ECO1123","คณิต–สถิติเบื้องต้นสำหรับ ศ.",3,4,"แกน • ด่านทดสอบเลข","-","✓","econ"),("ACC1103","การบัญชีการเงิน",3,4,"แกน econ","-","✓","econ"),("MTH1101","Calculus & Analytic Geometry I",3,4,"เลือกเสรี","-","✓","quant"),("RAM1111","English in Daily Life",3,2,"ศึกษาทั่วไป 1.1","-","✓","gen"),("RAM1103","Thai for Communication at Work",3,2,"ศึกษาทั่วไป 1.1","-","✓","gen"),("COS1101","Introduction to Computer Science",3,3,"เก็บเพิ่ม • พื้นฐาน CS","-","✓","quant")]),
 ("ปีที่ 1 / เทอม 2",[("ECO1122","หลักเศรษฐศาสตร์มหภาค",3,3,"แกน econ","ECO1123","✓","econ"),("MTH1102","Calculus & Analytic Geometry II",3,5,"เลือกเสรี","MTH1101","✓","quant"),("MTH2001","Fundamental Concepts in Math 1",3,7,"★ พิสูจน์ (ฐาน Real Analysis)","-","✓","quant"),("MTH2206","Matrix Theory & Linear Algebra 1",3,5,"โทคณิต","MTH1101","✓","quant"),("RAM1131","Digital Literacy",3,2,"ศึกษาทั่วไป 1.1","-","✓","gen"),("RAM1301","คุณธรรมคู่ความรู้",3,2,"ศึกษาทั่วไป 1.3","-","✓","gen"),("COS1103","Algorithms and Programming",3,5,"เก็บเพิ่ม • ปลดล็อก COS2102","COS1101","✓","quant")]),
 ("ปีที่ 2 / เทอม 1",[("ECO2121","ทฤษฎีจุลภาค 1",3,6,"แกน econ","ECO1121,1123","✓","econ"),("ECO2122","ทฤษฎีมหภาค 1",3,6,"แกน econ","ECO1122,1123","✓","econ"),("ECO2125","คณิตศาสตร์สำหรับเศรษฐศาสตร์",3,6,"แกน • วิชาสัญญาณ","ECO1123","✓","econ"),("ECO2129","เศรษฐกิจประเทศไทย",3,3,"แกน econ","ECO1121,1122","✓","econ"),("MTH2103","Calculus & Analytic Geometry III",3,6,"โทคณิต • หลายตัวแปร","MTH1102","✓","quant"),("STA3101","Probability Theory 1",3,7,"★ ฐานความน่าจะเป็น","MTH1102","✓","quant"),("RAM1112","ภาษาและวัฒนธรรมอังกฤษ",3,2,"ศึกษาทั่วไป 1.1","-","✓","gen")]),
 ("ปีที่ 2 / เทอม 2",[("ECO2123","ทฤษฎีจุลภาค 2",3,7,"แกน econ","ECO2121","✓","econ"),("ECO2124","ทฤษฎีมหภาค 2",3,7,"แกน econ","ECO2122","✓","econ"),("ECO2126","สถิติสำหรับเศรษฐศาสตร์",3,6,"แกน • วิชาสัญญาณ","STA2003 (ทางการ)","⚠️ ไม่มี STA2003 — เช็ก ECO1123","econ"),("MTH2104","Advanced Calculus 1",3,8,"โทคณิต • ปลดล็อก Real Analysis","MTH2103","✓","quant"),("MTH3204","Matrix Theory & Linear Algebra 2",3,7,"โทคณิต","MTH2206","✓","quant"),("RAM1141","ศาสตร์แห่งบุคลิกภาพ",3,2,"ศึกษาทั่วไป 1.1","-","✓","gen"),("ACC2202","Intermediate Accounting",3,6,"บัญชีเชิงลึก (ก่อนฝึกงาน)","ACC1103","⚠️ เช็ก ACC1101/1103","fin")]),
 ("ปีที่ 3 / เทอม 1",[("ECO3121","เศรษฐศาสตร์จุลภาคระดับสูง",3,8,"เชี่ยวชาญ","ECO2123,2125","✓","econ"),("ECO3123","เศรษฐมิติเบื้องต้น",3,7,"เชี่ยวชาญ • ดาวเด่น","ECO2125,2126","✓ (รอ ECO2126)","econ"),("ECO3320","เศรษฐศาสตร์การเงินและการธนาคาร",3,5,"เลือกกลุ่มอื่น–การเงิน","ECO2124","✓","fin"),("MTH2106","Differential Equations",3,6,"โทคณิต","MTH1102","✓","quant"),("STA3102","Probability Theory 2",3,8,"★ ปลดล็อก Theory of Stats","STA3101+MTH2103","✓","quant"),("ECO2130","English for Economics",3,3,"แกน econ","-","✓","econ"),("FIN2101","Business Finance",3,5,"เก็บเพิ่ม • บุพวิชา FIN3xxx","ACC1103","⚠️ เช็ก PR + สิทธิ์ข้ามคณะ","fin")]),
 ("ปีที่ 3 / เทอม 2",[("ECO3122","เศรษฐศาสตร์มหภาคระดับสูง",3,8,"เชี่ยวชาญ","ECO2124,2125","✓","econ"),("ECO3323","การบริหารความเสี่ยงองค์กร",3,5,"เลือกกลุ่มอื่น–การเงิน","ECO2121","✓","fin"),("ECO2128","ระเบียบวิธีวิจัยทางเศรษฐศาสตร์",3,5,"แกน • ต่อยอด research paper","ECO2123-2126","✓","econ"),("MTH4404","Introduction to Real Analysis 1",3,9,"เก็บเพิ่ม • สัญญาณ #1","MTH2104(+MTH2001)","✓","quant"),("ACC4209","Financial Reporting & Analysis",3,6,"วิเคราะห์งบ (core IB)","ACC2202","✓","fin"),("SDM2102","Foundations of Python Programming",3,4,"เก็บเพิ่ม • Python","-","✓","quant"),("RAM1312","Contemporary Culture & Digital Disruption",3,2,"⬅ ย้ายมา v10 • ลดโหลด","-","✓","gen")]),
 ("ปีที่ 4 / เทอม 1",[("ECO4124","Time-Series Econometrics",3,8,"เชี่ยวชาญ (เลือก)","ECO3123","✓","econ"),("ECO4125","Microeconometrics",3,8,"เชี่ยวชาญ • เช็ก causal inf.","ECO3123","✓","econ"),("ECO4326","Mathematical Finance",3,8,"เลือกกลุ่มอื่น–การเงิน","ECO3323","✓","fin"),("MTH3611","Mathematical Stochastic Models",3,8,"เก็บเพิ่ม • Stochastic","MTH2106","✓ PR คณิต","quant"),("FIN3211","Financial Management (Corp Finance)",3,6,"เก็บเพิ่ม • เนื้อ IB","FIN2101","⚠️ อาจต้อง FIN3208","fin"),("RAM1204","Mathematics & Statistics for Daily Life",3,3,"⬅ ย้ายมา v9","-","✓","gen"),("RAM1302","Politics and Law in Daily Life",3,2,"ศึกษาทั่วไป 1.3","-","✓","gen")]),
 ("ปีที่ 4 / เทอม 2",[("ECO4123","Game Theory (ทฤษฎีเกม)",3,7,"เชี่ยวชาญ (เลือก)","ECO2123,2125","✓","econ"),("ECO4327","Financial Risk Mgmt & Derivatives",3,7,"เลือกกลุ่มอื่น–การเงิน","ECO3320,3323","✓","fin"),("COS2102","Object-Oriented Programming (C++)",3,6,"เก็บเพิ่ม • C++ (quant dev)","COS1103","✓","quant"),("FIN3209","Securities Analysis",3,6,"เก็บเพิ่ม • valuation","FIN2101","✓","fin"),("FIN3205","Principles & Policies of Investment",3,5,"เก็บเพิ่ม • การลงทุน","FIN2101","✓","fin"),("RAM1211","Principles of Entrepreneurship",3,2,"ศึกษาทั่วไป 1.2","-","✓","gen"),("STA4101","Theory of Statistics 1",3,9,"★ Math Stats (MFE core)","STA3102","⬅ ย้ายจากปี3/ท2","quant")]),
]

wb=Workbook()
# ============ SHEET 1: PLAN ============
ws=wb.active; ws.title="แผนเรียน v10"; ws.sheet_view.showGridLines=False
H=["✅","รหัส","ชื่อวิชา","นก.","ยาก/10","หมวด / เหตุผล","บุพวิชา","สถานะ / หมายเหตุ"]; NC=8
ws.merge_cells("A1:H1"); c=ws["A1"]; c.value="แผนการเรียน ป.ตรี เศรษฐศาสตร์ ม.รามคำแหง — Quant+Finance/IB (v11 รวม GPAX)"; c.font=FT; c.fill=NAVY; c.alignment=C; ws.row_dimensions[1].height=24
ws.merge_cells("A2:H2"); c=ws["A2"]; c.value="ไฟล์รวม: ชีตนี้=แผนเรียน • ชีต 'GPAX'=คำนวณเกรด • ชีต 'เทียบ ป.โท Top-U'=โอกาสเรียนต่อ | ยาก/10: เขียว≤4 เหลือง5-6 ส้ม7-8 แดง9-10"; c.font=FNOTE; c.alignment=L; ws.row_dimensions[2].height=22
r=4
for j,h in enumerate(H,1):
    cc=ws.cell(r,j,h); cc.fill=BLUE; cc.font=FW; cc.alignment=C if j in(1,2,4,5) else L; cc.border=BD
ws.freeze_panes="A5"; r=5
for term,rows in PLAN:
    avg=round(sum(x[3] for x in rows)/len(rows),1)
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=NC-1)
    tc=ws.cell(r,1,"%s   (ยากเฉลี่ย %s/10)"%(term,avg)); tc.fill=PANEL; tc.font=FB; tc.alignment=L; tc.border=BD
    for j in range(2,NC): ws.cell(r,j).fill=PANEL; ws.cell(r,j).border=BD
    cc=ws.cell(r,NC,"21 นก."); cc.fill=PANEL; cc.font=FB; cc.alignment=C; cc.border=BD; r+=1
    for code,name,cr,diff,reason,pre,status,cat in rows:
        for j,v in enumerate(["",code,name,cr,diff,reason,pre,status],1):
            cell=ws.cell(r,j,v); cell.fill=FILL[cat]; cell.font=F; cell.border=BD
            cell.alignment=C if j in(1,2,4,5) else L
        ws.cell(r,5).fill=dfill(diff); ws.cell(r,5).font=FB; r+=1
for col,w in zip("ABCDEFGH",[4,9,40,5,8,36,22,30]): ws.column_dimensions[col].width=w

# ============ SHEET 2: GPAX ============
g=wb.create_sheet("GPAX"); g.sheet_view.showGridLines=False
courses=[]
for term,rows in PLAN:
    m=re.match(r'ปีที่\s*(\d+)\s*/\s*เทอม\s*(\d+)',term); ts="%s/%s"%(m.group(1),m.group(2))
    for code,name,cr,*_ in rows: courses.append((ts,code,name,cr))
FIRST=10; NX=5; LAST=FIRST+len(courses)+NX-1
Dr="$D$%d:$D$%d"%(FIRST,LAST); Er="$E$%d:$E$%d"%(FIRST,LAST); Gr="$G$%d:$G$%d"%(FIRST,LAST); Ar="$A$%d:$A$%d"%(FIRST,LAST)
g.merge_cells("A1:G1"); c=g["A1"]; c.value="🎓 คำนวณ GPAX อัตโนมัติ — เลือกเกรดในช่องสีเหลือง"; c.font=FT; c.fill=NAVY; c.alignment=C; g.row_dimensions[1].height=24
g.merge_cells("A2:G2"); g["A2"]="A=4.0 B+=3.5 B=3.0 C+=2.5 C=2.0 D+=1.5 D=1.0 F=0.0 • GPAX=Σ(แต้ม×นก.)÷Σนก.ที่มีเกรด (F นับ 0 แต่รวมหน่วยกิต)"; g["A2"].font=FNOTE; g["A2"].alignment=L
g.merge_cells("A4:G4"); c=g["A4"]; c.value="📊 สรุปผล (คำนวณอัตโนมัติ)"; c.font=FW; c.fill=BLUE; c.alignment=L
def lab(cell,txt): g[cell]=txt; g[cell].font=FB; g[cell].fill=PANEL; g[cell].alignment=L; g[cell].border=BD
g.merge_cells("A5:B5"); lab("A5","GPAX สะสม")
g.merge_cells("C5:D5"); g["C5"]="=IF($G$5=0,\"\",SUM(%s)/$G$5)"%Gr; g["C5"].font=FBIG; g["C5"].fill=GPAXC; g["C5"].alignment=C; g["C5"].number_format="0.00"; g["C5"].border=BD
g.merge_cells("E5:F5"); lab("E5","หน่วยกิตที่มีเกรดแล้ว"); g["G5"]="=SUMPRODUCT((%s<>\"\")*%s)"%(Er,Dr); g["G5"].font=FB; g["G5"].alignment=C; g["G5"].border=BD; g["G5"].number_format="0"; g.row_dimensions[5].height=30
g.merge_cells("A6:B6"); lab("A6","เทียบเกียรตินิยม*")
g.merge_cells("C6:D6"); g["C6"]="=IF($C$5=\"\",\"-\",IF($C$5>=3.8,\"อันดับ 1 เหรียญทอง 🥇\",IF($C$5>=3.6,\"อันดับ 1 🏆\",IF($C$5>=3.3,\"อันดับ 2 🥈\",\"ไม่เข้าเกณฑ์\"))))"; g["C6"].font=FB; g["C6"].alignment=C; g["C6"].border=BD
g.merge_cells("E6:F6"); lab("E6","หน่วยกิตรวมทั้งหมด"); g["G6"]="=SUM(%s)"%Dr; g["G6"].font=FB; g["G6"].alignment=C; g["G6"].border=BD; g["G6"].number_format="0"
g.merge_cells("A7:B7"); lab("A7","🎯 ตั้งเป้า GPAX (กรอก)")
g.merge_cells("C7:D7"); g["C7"]=3.80; g["C7"].font=FB; g["C7"].fill=INPUT; g["C7"].alignment=C; g["C7"].border=BD; g["C7"].number_format="0.00"
g.merge_cells("E7:F7"); lab("E7","ต้องได้เฉลี่ยอีก (เต็ม 4)"); g["G7"]="=IF(($G$6-$G$5)<=0,\"-\",IF($C$7=\"\",\"\",($C$7*$G$6-SUM(%s))/($G$6-$G$5)))"%Gr; g["G7"].font=FB; g["G7"].alignment=C; g["G7"].border=BD; g["G7"].number_format="0.00"
g.merge_cells("A8:G8"); g["A8"]="เกณฑ์เกียรตินิยม: เหรียญทอง ≥3.80 | อันดับ1 = 3.60–3.79 | อันดับ2 = 3.30–3.59 (ต่ำกว่า 3.30 = ไม่เข้าเกณฑ์) • ดูชีต 'เทียบ ป.โท Top-U' ประกอบ"; g["A8"].font=FNOTE; g["A8"].alignment=L; g["A8"].border=BD; g.row_dimensions[8].height=22
hd=["ปี/เทอม","รหัส","ชื่อวิชา","นก.","เกรด","แต้ม","แต้ม×นก."]
for j,h in enumerate(hd,1):
    cc=g.cell(9,j,h); cc.fill=BLUE; cc.font=FW; cc.alignment=C; cc.border=BD
g.freeze_panes="A10"; r=FIRST
for ts,code,name,cr in courses:
    g.cell(r,1,ts).alignment=C; g.cell(r,2,code).alignment=C; g.cell(r,3,name); g.cell(r,4,cr).alignment=C
    g.cell(r,5).fill=INPUT; g.cell(r,5).alignment=C
    g.cell(r,6,"=IFERROR(VLOOKUP($E%d,Scale!$A$2:$B$9,2,FALSE),\"\")"%r).alignment=C
    g.cell(r,7,"=IF($E%d=\"\",0,$F%d*$D%d)"%(r,r,r)).alignment=C
    g.cell(r,6).number_format="0.0"; g.cell(r,7).number_format="0.0"
    for j in range(1,8): g.cell(r,j).font=F; g.cell(r,j).border=BD
    r+=1
for _ in range(NX):
    g.cell(r,5).fill=INPUT; g.cell(r,5).alignment=C
    g.cell(r,6,"=IFERROR(VLOOKUP($E%d,Scale!$A$2:$B$9,2,FALSE),\"\")"%r).alignment=C
    g.cell(r,7,"=IF($E%d=\"\",0,$F%d*$D%d)"%(r,r,r)).alignment=C
    for j in range(1,8): g.cell(r,j).font=F; g.cell(r,j).border=BD
    r+=1
g.cell(r,3,"รวม").font=FB; g.cell(r,3).alignment=Alignment(horizontal="right")
g.cell(r,4,"=SUM(%s)"%Dr).font=FB; g.cell(r,4).alignment=C; g.cell(r,4).number_format="0"
g.cell(r,7,"=SUM(%s)"%Gr).font=FB; g.cell(r,7).alignment=C; g.cell(r,7).number_format="0.0"
for j in range(1,8): g.cell(r,j).fill=PANEL; g.cell(r,j).border=BD
dv=DataValidation(type="list",formula1='"A,B+,B,C+,C,D+,D,F"',allow_blank=True); g.add_data_validation(dv); dv.add("E%d:E%d"%(FIRST,LAST))
gz="E%d:E%d"%(FIRST,LAST)
for grd,fl in [("A",D_GRN),("B+",D_GRN),("D",D_ORG),("D+",D_ORG),("F",D_RED)]:
    g.conditional_formatting.add(gz,CellIsRule(operator="equal",formula=['"%s"'%grd],fill=fl))
g.merge_cells("I4:K4"); h=g["I4"]; h.value="GPA รายเทอม"; h.fill=BLUE; h.font=FW; h.alignment=C; h.border=BD
g.cell(5,9,"เทอม").font=FB; g.cell(5,10,"GPA").font=FB; g.cell(5,11,"นก.").font=FB
for j in(9,10,11): g.cell(5,j).fill=PANEL; g.cell(5,j).alignment=C; g.cell(5,j).border=BD
for i,t in enumerate(["1/1","1/2","2/1","2/2","3/1","3/2","4/1","4/2"]):
    rr=6+i; g.cell(rr,9,t).alignment=C; g.cell(rr,9).font=F
    crf="SUMPRODUCT((%s=$I%d)*(%s<>\"\")*%s)"%(Ar,rr,Er,Dr); qf="SUMPRODUCT((%s=$I%d)*%s)"%(Ar,rr,Gr)
    g.cell(rr,10,"=IF(%s=0,\"\",%s/%s)"%(crf,qf,crf)).number_format="0.00"; g.cell(rr,10).font=FB; g.cell(rr,10).alignment=C
    g.cell(rr,11,"=%s"%crf).number_format="0"; g.cell(rr,11).font=F; g.cell(rr,11).alignment=C
    for j in(9,10,11): g.cell(rr,j).border=BD
g.conditional_formatting.add("J6:J13",ColorScaleRule(start_type="num",start_value=0,start_color="F8696B",mid_type="num",mid_value=2,mid_color="FFEB84",end_type="num",end_value=4,end_color="63BE7B"))
for col,w in zip("ABCDEFG",[8,9,40,5,8,7,10]): g.column_dimensions[col].width=w
for col,w in zip("IJK",[8,8,7]): g.column_dimensions[col].width=w

# ============ SHEET 3: COMPARE ป.โท Top-U ============
cm=wb.create_sheet("เทียบ ป.โท Top-U"); cm.sheet_view.showGridLines=False; NCc=6
def merge(row,txt,fill,font,al=L,h=None):
    cm.merge_cells(start_row=row,start_column=1,end_row=row,end_column=NCc)
    c=cm.cell(row,1,txt); c.font=font; c.alignment=al
    if fill: c.fill=fill
    if h: cm.row_dimensions[row].height=h
merge(1,"เทียบเกณฑ์ GPAX กับการยื่น ป.โท Top-20 (QS) — โอกาส & ความเสี่ยง",NAVY,FT,C,24)
merge(2,"อ้างอิงเกณฑ์รับเข้าทั่วไป (2026): US ป.โทส่วนใหญ่ขั้นต่ำ 3.0 (70% ขอ 3.0–3.5) แข่งขัน 3.5–3.7 • elite quant ค่าเฉลี่ยที่รับ ~3.7–3.8 (MIT MFin 3.7) • UK: 2:1 ≈ GPA 3.3–3.6, First ≈ 3.7–4.0 (LSE: 3.5=2:1; Oxford strong-2:1=3.5, First=3.7) • LSE/Oxford ที่รับจริงส่วนใหญ่ได้ First",None,FNOTE,L,46)
# main mapping table
r=4; merge(r,"① แมปเกณฑ์เกียรตินิยมของคุณ → โอกาสยื่น Top-20 (แกน GPA)",SEC,FSEC); r+=1
heads=["เกียรตินิยม","GPAX","เทียบ UK/US","โอกาส Top-20 (เชิง GPA)","ความเสี่ยง",""]
for j,h in enumerate(heads[:5],1):
    c=cm.cell(r,j,h); c.fill=BLUE; c.font=FW; c.alignment=C; c.border=BD
cm.merge_cells(start_row=r,start_column=5,end_row=r,end_column=6); r+=1
rows=[
 ("เหรียญทอง","≥ 3.80","First (เกิน) • เกิน MIT MFin avg","แข่งขันได้แทบทุกโปรแกรม Top-20","🟢 ต่ำ","low"),
 ("อันดับ 1","3.60–3.79","First / high 2:1 • 3.7=avg MIT MFin","แข่งได้ดี-กลาง (Top-5 = กลางๆ)","🟡 กลาง","med"),
 ("อันดับ 2","3.30–3.59","2:1 (ผ่าน min LSE/Oxford/US)","ผ่านขั้นต่ำ แต่ต่ำกว่าค่าเฉลี่ยที่รับ","🟠 สูง (โดยเฉพาะ Top-5)","high"),
 ("(ต่ำกว่าเกณฑ์)","< 3.30","ต่ำกว่า 2:1","ต่ำกว่าขั้นต่ำหลายที่","🔴 สูงมาก","vhigh"),
]
for a,b,c2,d,e,risk in rows:
    for j,v in enumerate([a,b,c2,d,e],1):
        cell=cm.cell(r,j,v); cell.font=FB if j==1 else F; cell.border=BD; cell.alignment=L if j in(3,4) else C
    cm.merge_cells(start_row=r,start_column=5,end_row=r,end_column=6)
    cm.cell(r,5).fill=RISK[risk]; cm.cell(r,5).alignment=C; cm.cell(r,5).font=FB; cm.cell(r,6).fill=RISK[risk]; cm.cell(r,6).border=BD
    cm.row_dimensions[r].height=28; r+=1
r+=1
merge(r,"② ⚠️ ปัจจัยที่สำคัญกว่า GPAX (ต้องรู้!)",SEC,FSEC); r+=1
for ln in [
 "• ม.รามคำแหง = 'non-target / open university' → adcom อ่าน GPA ตามบริบทสถาบัน → GPAX ต้อง 'สูงกว่าค่าเฉลี่ย' เพื่อชดเชย brand discount",
 "• ตัวชี้วัดที่ขยับผลจริง: ① GRE Quant สูง (168–170) ② เกรด A วิชาสัญญาณ (Real Analysis, Probability, Theory of Stats, Econometrics) ③ research/writing sample ④ จดหมายแนะนำแรง ⑤ ประสบการณ์/predoc/RA",
 "• สรุป: GPAX สูง = 'จำเป็นแต่ไม่พอ' — ผ่านด่านคัดกรอง แต่ต้องมีแพ็กเกจครบจึงชนะคู่แข่งจาก target school",
 "• เคล็ด: ถึง GPAX รวมจะโดน gen-ed ฉุด ให้รักษาเกรดวิชาสัญญาณ (สีเขียวในแผน) ให้ A ไว้ — adcom quant ดูตรงนี้ก่อน GPAX รวม",
]:
    merge(r,ln,None,F,LT,28); cm.cell(r,1).border=BD; r+=1
r+=1
merge(r,"③ สรุปความเสี่ยงเชิงปฏิบัติ (เมื่อรวม brand ม.ราม)",SEC,FSEC); r+=1
for ln,fl in [
 ("เหรียญทอง (≥3.80): Top-20 = แข่งได้จริง (ต้องมี GRE+research+LoR ครบ) • Top-5 (MIT/Stanford/Harvard/LSE) = เอื้อมได้แต่ต้องโปรไฟล์เต็ม","low"),
 ("อันดับ 1 (3.60–3.79): Top-20 กลาง-ล่าง = โอกาสดี • Top-5 = เสี่ยง ต้องชดเชยหนักด้วย GRE/research","med"),
 ("อันดับ 2 (3.30–3.59): Top-20 = เสี่ยงสูง • ควรเล็ง Top-20 อันดับท้าย + โปรแกรมที่ fit + ชดเชยด้วย GRE/research มาก หรือไต่ผ่าน predoc/ทำงานก่อน","high"),
]:
    merge(r,ln,RISK[fl],FB,LT,30); cm.cell(r,1).border=BD; r+=1
r+=1
merge(r,"แหล่งอ้างอิง: research.com (Finance/Econ Master's 2026) • LSE MSc Finance & Economics • Oxford International Qualifications • QuantNet/MIT MFin • เกณฑ์เปลี่ยนได้ ตรวจหน้าโปรแกรมจริงก่อนสมัคร",None,FNOTE,L,30)
for col,w in zip("ABCDEF",[16,11,30,34,16,6]): cm.column_dimensions[col].width=w

# ============ SHEET 4: Scale (hidden) ============
sc=wb.create_sheet("Scale"); sc["A1"]="เกรด"; sc["B1"]="แต้ม"
for i,(gr,p) in enumerate([("A",4),("B+",3.5),("B",3),("C+",2.5),("C",2),("D+",1.5),("D",1),("F",0)],2): sc.cell(i,1,gr); sc.cell(i,2,p)
sc.sheet_state="hidden"

wb.active=0
try: wb.calculation.fullCalcOnLoad=True
except Exception: pass
wb.save(OUT); print("SAVED",OUT,"sheets:",wb.sheetnames)
